"""文档生成工具"""
from pathlib import Path
from .base import BaseTool, ToolResult, tool_registry


class CreateDocxTool(BaseTool):
    name = "create_docx"
    description = "创建 Word 文档（.docx）"
    parameters_schema = {
        "filename": {"type": "string", "description": "文件名（不含扩展名）"},
        "content": {"type": "string", "description": "文档内容（支持 Markdown 格式）"},
        "title": {"type": "string", "description": "文档标题", "optional": True}
    }

    def __init__(self, workspace: str):
        self.workspace = Path(workspace)

    def execute(self, filename: str, content: str, title: str = "") -> ToolResult:
        try:
            from docx import Document
            from docx.shared import Pt

            doc = Document()

            if title:
                doc.add_heading(title, 0)

            # 解析 Markdown
            for line in content.split("\n"):
                line = line.strip()
                if not line:
                    continue

                if line.startswith("# "):
                    doc.add_heading(line[2:], 1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], 2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:], 3)
                elif line.startswith("- ") or line.startswith("* "):
                    doc.add_paragraph(line[2:], style="List Bullet")
                elif len(line) > 2 and line[0].isdigit() and line[1] in ".)":
                    doc.add_paragraph(line[2:].strip(), style="List Number")
                else:
                    doc.add_paragraph(line)

            filepath = self.workspace / f"{filename}.docx"
            doc.save(str(filepath))

            return ToolResult(True, f"已创建: {filename}.docx", artifacts=[str(filepath)])
        except ImportError:
            return ToolResult(False, "", "缺少 python-docx 库")
        except Exception as e:
            return ToolResult(False, "", str(e))


class CreateXlsxTool(BaseTool):
    name = "create_xlsx"
    description = "创建 Excel 表格（.xlsx）"
    parameters_schema = {
        "filename": {"type": "string", "description": "文件名（不含扩展名）"},
        "data": {"type": "string", "description": "CSV 格式数据（逗号分隔列，换行分隔行）"},
        "sheet_name": {"type": "string", "description": "工作表名称", "optional": True}
    }

    def __init__(self, workspace: str):
        self.workspace = Path(workspace)

    def execute(self, filename: str, data: str, sheet_name: str = "Sheet1") -> ToolResult:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name

            for row_idx, row in enumerate(data.strip().split("\n"), 1):
                for col_idx, cell in enumerate(row.split(","), 1):
                    ws.cell(row=row_idx, column=col_idx, value=cell.strip())

            # 调整列宽
            for col in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col)].width = 15

            filepath = self.workspace / f"{filename}.xlsx"
            wb.save(str(filepath))

            return ToolResult(True, f"已创建: {filename}.xlsx", artifacts=[str(filepath)])
        except ImportError:
            return ToolResult(False, "", "缺少 openpyxl 库")
        except Exception as e:
            return ToolResult(False, "", str(e))


class CreatePptxTool(BaseTool):
    name = "create_pptx"
    description = "创建 PowerPoint 演示文稿（.pptx）"
    parameters_schema = {
        "filename": {"type": "string", "description": "文件名（不含扩展名）"},
        "slides": {"type": "string", "description": "幻灯片内容（用 --- 分隔每张，第一行为标题）"}
    }

    def __init__(self, workspace: str):
        self.workspace = Path(workspace)

    def execute(self, filename: str, slides: str) -> ToolResult:
        try:
            from pptx import Presentation

            prs = Presentation()

            for slide_content in slides.split("---"):
                lines = [l.strip() for l in slide_content.strip().split("\n") if l.strip()]
                if not lines:
                    continue

                title = lines[0].lstrip("#").strip()
                body = "\n".join(lines[1:]) if len(lines) > 1 else ""

                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                slide.shapes.title.text = title

                if body and len(slide.placeholders) > 1:
                    slide.placeholders[1].text = body

            filepath = self.workspace / f"{filename}.pptx"
            prs.save(str(filepath))

            return ToolResult(
                True,
                f"已创建: {filename}.pptx ({len(prs.slides)} 张幻灯片)",
                artifacts=[str(filepath)]
            )
        except ImportError:
            return ToolResult(False, "", "缺少 python-pptx 库")
        except Exception as e:
            return ToolResult(False, "", str(e))


def register_document_tools(workspace: str):
    """注册文档工具"""
    tool_registry.register(CreateDocxTool(workspace))
    tool_registry.register(CreateXlsxTool(workspace))
    tool_registry.register(CreatePptxTool(workspace))
